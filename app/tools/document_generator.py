from pathlib import Path
from docx import Document

BASE_DIR = Path(__file__).resolve().parents[2]
OUTPUT_DIR = (BASE_DIR / "data" / "output").resolve()


def create_inspection_note(
    equipment: str,
    inspection_date: str,
    findings: list[str],
    recommendations: list[str],
    output_filename: str = "inspection_review_note.docx",
) -> str:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    output_path = (OUTPUT_DIR / output_filename).resolve()

    if OUTPUT_DIR not in output_path.parents:
        raise ValueError("Output path is outside the allowed directory.")

    doc = Document()

    doc.add_heading("Inspection Review Note", level=1)
    doc.add_paragraph("Sovereign AI Workbench — Local Demonstration")

    doc.add_heading("Equipment", level=2)
    doc.add_paragraph(equipment)

    doc.add_heading("Inspection Date", level=2)
    doc.add_paragraph(inspection_date)

    doc.add_heading("Key Findings", level=2)
    for finding in findings:
        doc.add_paragraph(finding, style="List Bullet")

    doc.add_heading("Recommendations", level=2)
    for recommendation in recommendations:
        doc.add_paragraph(recommendation, style="List Bullet")

    doc.add_heading("Review Status", level=2)
    doc.add_paragraph(
        "Generated for technical review. Final engineering and approval "
        "decisions remain with authorized personnel."
    )

    doc.add_paragraph()
    doc.add_paragraph(
        "Data sovereignty: Generated locally using the Sovereign AI Workbench prototype."
    )

    doc.save(output_path)

    return str(output_path)


if __name__ == "__main__":
    path = create_inspection_note(
        equipment="Heat Exchanger HX-101",
        inspection_date="2026-08-20",
        findings=[
            "Minor surface corrosion was observed near the outlet nozzle.",
            "Measured wall thickness at inspection point P3 was 8.4 mm.",
            "Nominal wall thickness is 10.0 mm.",
            "No visible leakage was observed.",
        ],
        recommendations=[
            "Continue periodic thickness monitoring.",
            "Perform detailed corrosion assessment during the next planned maintenance window.",
            "Review inspection history before deciding on further maintenance.",
        ],
    )

    print(f"DOCX CREATED: {path}")

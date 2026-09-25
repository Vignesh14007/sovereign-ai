from pathlib import Path
from typing import Optional

import pymupdf
from docx import Document


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".pdf",
    ".docx",
}


def extract_text_from_txt(file_path: Path) -> str:
    """Extract text from a plain-text file."""
    return file_path.read_text(encoding="utf-8", errors="ignore")


def extract_text_from_pdf(file_path: Path) -> str:
    """Extract text from a PDF using PyMuPDF."""

    pages = []

    with pymupdf.open(file_path) as pdf:

        for page in pdf:
            text = page.get_text("text")

            if text:
                pages.append(text)

    return "\n\n".join(pages)


def extract_text_from_docx(file_path: Path) -> str:
    """Extract paragraphs and tables from a DOCX file."""

    document = Document(file_path)

    parts = []

    # Normal paragraphs
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            cells = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    cells.append(text)

            if cells:
                parts.append(" | ".join(cells))

    return "\n\n".join(parts)


def extract_text(file_path: str | Path) -> str:
    """
    Extract text from a supported document.

    Supported formats:
        TXT
        PDF
        DOCX
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))

        raise ValueError(
            f"Unsupported file type: {extension}. "
            f"Supported types: {supported}"
        )

    if extension == ".txt":
        text = extract_text_from_txt(path)

    elif extension == ".pdf":
        text = extract_text_from_pdf(path)

    elif extension == ".docx":
        text = extract_text_from_docx(path)

    else:
        raise ValueError(f"Unsupported file type: {extension}")

    text = text.strip()

    if not text:
        raise ValueError(
            f"No text could be extracted from: {path.name}"
        )

    return text


def get_file_metadata(file_path: str | Path) -> dict:
    """
    Return basic metadata for an uploaded document.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    return {
        "name": path.name,
        "extension": path.suffix.lower(),
        "size": path.stat().st_size,
    }
